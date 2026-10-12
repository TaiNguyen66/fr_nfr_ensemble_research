"""Tạo workbook sàng lọc tài liệu (review_protocol.md §4–§7) từ kết quả của frnfr-metadata-fetcher.

Chạy (từ gốc repo), sau `cli.mjs screen`:
    python scripts/build_screening_workbook.py

Đọc literature/search/01_all_records.csv và 02_after_screening_v1.csv, ghi
literature/review/NCKH_FRNFR_Literature.xlsx gồm các tab:
    0_HuongDan        tiêu chí và mã quyết định cho từng vòng
    R1_TieuDe_TomTat  vòng 1: mọi bản ghi cần người đọc tiêu đề + tóm tắt, đã chia người
    R2_DuDieuKien     vòng 2: kiểm tra điều kiện (chép các dòng KEEP/MAYBE của R1 sang)
    R3_ToanVan_QA     vòng 3: đọc toàn văn, trích xuất, chấm chất lượng
    GD3_Gap           giai đoạn 3: bảng khoảng trống → RQ

Đưa lên Drive rồi "File → Lưu dưới dạng Google Trang tính" để hai người cùng sửa.
Gợi ý tự động chỉ để tham khảo; quyết định là của người sàng lọc.
"""

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "literature" / "search"
OUT = ROOT / "literature" / "review" / "NCKH_FRNFR_Literature.xlsx"

CORE = {"Q1", "Q1b", "Q2", "Q3", "Q4", "Q5"}
T_CAP = 40          # số bài nền tảng T1/T2 tối đa đưa vào R1
OVERLAP_EVERY = 10  # mỗi 10 bài có 1 bài cả hai cùng sàng lọc (đo độ đồng thuận)
A, B = "Hoang Anh", "Phat Tai"

# Paper nền không có trong kết quả tìm kiếm (trước 2017 hoặc ngoài chuỗi), thêm tay
SEEDS = [
    ("S01", "Automated classification of non-functional requirements", "Cleland-Huang et al.", "2007",
     "Requirements Engineering 12(2)", "", "Nguồn gốc bộ PROMISE NFR"),
    ("S02", "Software Engineering Repositories: Expanding the PROMISE Database", "Lima et al. (cần xác minh)", "2019",
     "SBES 2019 (cần xác minh)", "", "Nguồn gốc PROMISE_exp: xác minh trước khi trích"),
    ("S03", "Promise+: expandindo a base de dados de requisitos de software Promise_exp", "(xem bài)", "2024",
     "SBES 2024", "", "Mô tả 444 FR / 525 NFR, cột ProjectID"),
    ("S04", "Natural Language Processing for Requirements Engineering: A Systematic Mapping Study", "Zhao et al.",
     "2021", "ACM Computing Surveys 54(3)", "10.1145/3444689", "Bối cảnh NLP4RE"),
    ("S05", "Inference for the Generalization Error", "Nadeau & Bengio", "2003", "Machine Learning 52", "",
     "Corrected resampled t-test (phương pháp)"),
]

DECISIONS = ["KEEP", "MAYBE", "DROP"]
CODES = "EC-S, EC-T, EC-E, EC-D, EC-Y, EC-G, EC-F"

GUIDE = [
    ("Mục đích", "Sàng lọc tài liệu theo review_protocol.md trong repo GitHub. Mỗi vòng có tab riêng; chỉ những dòng được giữ mới sang vòng sau."),
    ("Quyết định", "KEEP = giữ · MAYBE = chưa chắc, cho sang vòng sau · DROP = loại (ghi mã lý do)."),
    ("Mã lý do khi DROP", "EC-S ngoài phạm vi (không phải phân loại yêu cầu) · EC-T sai nhiệm vụ / loại văn bản (app review, issue, legal…) hoặc không dùng NLP/ML · "
     "EC-E không có đánh giá định lượng · EC-D trùng bài khác (ghi ID bài giữ lại) · EC-Y trước 2017 (trừ paper nền) hoặc không phải tiếng Anh · "
     "EC-G không phải bài nghiên cứu (dataset, slide, luận văn chưa bình duyệt) · EC-F không lấy được toàn văn (chỉ dùng ở R2/R3)"),
    ("R1 – Tiêu đề + tóm tắt", "Cột 'Phân công' cho biết ai sàng lọc. Dòng ghi 'Cả hai' thì mỗi người điền cột của mình độc lập, KHÔNG nhìn cột người kia (để đo độ đồng thuận). "
     "KEEP nếu: (IC1) bài tự động phân loại yêu cầu phần mềm (FR/NFR, loại NFR hoặc loại yêu cầu) bằng NLP/ML, VÀ (IC2) có đánh giá định lượng. "
     "Bài nền tảng (leakage, cross-validation, so sánh classifier, survey NLP4RE) giữ nếu dùng được cho Related Work hoặc phương pháp. Không chắc → MAYBE."),
    ("R1 – Gợi ý tự động", "Cột 'Gợi ý tự động' do luật từ khóa sinh ra, chỉ để tham khảo. Sai thì cứ quyết theo đánh giá của mình."),
    ("R2 – Đủ điều kiện", "Chép các dòng KEEP/MAYBE của R1 sang. Mở toàn văn, đọc lướt ~5 phút/bài, điền Y/N cho E1–E5. "
     "Bắt buộc E2, E4, E5 = Y để sang R3. Mục tiêu: 30–40 bài sang R3."),
    ("R3 – Toàn văn + QA", "Đọc toàn văn, điền các cột trích xuất và chấm Q1–Q5 (0/1/2). Tổng ≥ 6 = đưa vào tổng hợp chính; < 6 = LOW-CONFIDENCE (vẫn giữ, ghi rõ). "
     "Mục tiêu chốt: 15–20 bài (≥ 12 bài nghiên cứu thực nghiệm + 3–5 bài nền tảng)."),
    ("GĐ3 – Gap", "Mỗi dòng là một khoảng trống: loại gap, bằng chứng (ID bài + số bài / tổng), vì sao quan trọng, có làm được trong nguồn lực không (review_protocol.md §8), RQ đề xuất."),
    ("Sau khi sửa", "Không xóa dòng, không đổi cột ID. Muốn thêm bài ngoài danh sách: thêm dòng cuối tab R1 với ID dạng M01, M02… và ghi nguồn ở cột Ghi chú."),
]

R2_COLS = ["ID", "Tiêu đề", "Năm", "DOI / link", "Người đọc",
           "E1 Dataset có tên và truy cập được (Y/N)", "E2 Có kết quả cho FR/NFR hoặc loại yêu cầu (Y/N)",
           "E3 Kỹ thuật (ML cổ điển / ensemble / DL / LLM)", "E4 Mô tả cách chia dữ liệu / đánh giá (Y/N)",
           "E5 Có toàn văn (Y/N)", "Liên quan D1 ensemble (Y/N)", "Liên quan D2 tiền xử lý / biểu diễn (Y/N)",
           "Liên quan D3 đánh giá / leakage / tổng quát hóa (Y/N)", "Quyết định R2", "Mã lý do", "Ghi chú"]
R3_COLS = ["ID", "Tiêu đề", "Tác giả", "Năm", "Venue", "DOI", "Người đọc", "Nhóm (CORE-ML / CORE-ENSEMBLE / EVALUATION / BACKGROUND)",
           "Dataset (+version)", "Kích thước", "Nhiệm vụ (binary / multi-class / multi-label)", "Biểu diễn", "Model", "Tiền xử lý",
           "Cách chia dữ liệu", "Kiểm soát trùng lặp / leakage", "Metric", "Loại F1 (macro / weighted / per-class)",
           "Kết quả tốt nhất (kèm điều kiện)", "Kiểm định thống kê", "Code / data công khai", "So được với mình? (Yes / Partial / No)",
           "Hạn chế tác giả nêu", "Hạn chế mình thấy", "Ý chính cho đề tài",
           "Q1 Mục tiêu rõ (0-2)", "Q2 Dataset mô tả và truy cập được (0-2)", "Q3 Đánh giá chặt chẽ (0-2)",
           "Q4 Kết quả đủ và đúng metric (0-2)", "Q5 Hạn chế + tái lập (0-2)", "QA tổng", "Vị trí bằng chứng (section / table)", "Quyết định R3"]
GAP_COLS = ["Gap ID", "Loại gap (phương pháp đánh giá / so sánh / dữ liệu / tái lập / kỹ thuật)", "Mô tả khoảng trống",
            "Bằng chứng (ID bài)", "Số bài thể hiện / tổng bài", "Vì sao quan trọng", "Làm được trong nguồn lực? (Y/N + lý do)",
            "RQ đề xuất", "Liên quan RQ ứng viên (RQ1 / RQ2 / RQ3 / mới)", "Quyết định nhóm", "Ghi chú"]


def streams(x):
    return {s.strip() for s in x["search_stream"].split(";") if s.strip()}


def style_header(ws, widths):
    fill = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "C2"


def main():
    sys.stdout.reconfigure(encoding="utf-8")  # console Windows mặc định cp1252
    rows = list(csv.DictReader((SRC / "02_after_screening_v1.csv").open(encoding="utf-8-sig")))
    base = {x["id"]: x for x in csv.DictReader((SRC / "01_all_records.csv").open(encoding="utf-8-sig"))}
    cites = lambda x: int(base[x["id"]].get("cited_by_count") or 0)
    rank = {"P1": 0, "P2": 1, "P3": 2, "": 3}

    core = [x for x in rows if streams(x) & CORE and x["decision"] in ("INCLUDE", "UNSURE", "BACKGROUND")]
    core.sort(key=lambda x: (rank.get(x["v2_priority"], 3), -cites(x)))
    found = [x for x in rows if not streams(x) & CORE and x["decision"] != "EXCLUDE"
             and re.search(r"leak|duplicat|cross-validation|compar\w* (classifiers|learning algorithms)|t-test|reproducib",
                           x["title"], re.I)]
    found.sort(key=cites, reverse=True)
    pool = [("Chính", x) for x in core] + [("Nền tảng", x) for x in found[:T_CAP]]

    wb = Workbook()
    g = wb.active
    g.title = "0_HuongDan"
    g.append(["Mục", "Nội dung"])
    for row in GUIDE:
        g.append(list(row))
    style_header(g, [24, 140])
    for r in g.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    r1 = wb.create_sheet("R1_TieuDe_TomTat")
    cols = ["ID", "Phân công", "Loại", "Tiêu đề", "Năm", "Venue", "DOI / link", "Trích dẫn", "Gợi ý tự động",
            "Lý do gợi ý", "Ưu tiên", f"{A}: quyết định", f"{B}: quyết định", "Quyết định R1", "Mã lý do", "Ghi chú", "Tóm tắt"]
    r1.append(cols)
    # Phat Tai (chủ trì literature) ~60%, Hoang Anh ~40% vì còn mạch dữ liệu song song; xen kẽ để cả hai đều có bài P1
    pattern = [B, A, B, A, B]
    n = solo = 0
    for kind, x in pool:
        n += 1
        if n % OVERLAP_EVERY == 0:
            who = "Cả hai"
        else:
            who = pattern[solo % len(pattern)]
            solo += 1
        link = f"https://doi.org/{x['doi']}" if x["doi"] else x["url"]
        r1.append([x["id"], who, kind, x["title"], x["year"], x["venue"][:60], link, cites(x), x["decision"],
                   x["reason"], x["v2_priority"] or "-", "", "", "", "", "", x["abstract"][:1500]])
    for sid, title, auth, year, venue, doi, why in SEEDS:
        r1.append([sid, "Cả hai", "Paper nền (thêm tay)", title, year, venue, f"https://doi.org/{doi}" if doi else "",
                   "", "SEED", why, "P1", "", "", "", "", auth, ""])
    style_header(r1, [8, 11, 10, 60, 6, 28, 26, 8, 11, 34, 7, 14, 14, 13, 9, 24, 80])
    last = r1.max_row
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    r1.add_data_validation(dv)
    dv.add(f"L2:N{last}")

    for name, cols_, widths in [("R2_DuDieuKien", R2_COLS, [8, 50, 6, 26, 12] + [14] * 11),
                                ("R3_ToanVan_QA", R3_COLS, [8, 45, 18, 6, 18, 20, 12] + [16] * 26),
                                ("GD3_Gap", GAP_COLS, [8, 24, 50, 20, 12, 36, 24, 40, 16, 14, 24])]:
        ws = wb.create_sheet(name)
        ws.append(cols_)
        style_header(ws, widths)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    by_who = {}
    for row in r1.iter_rows(min_row=2, values_only=True):
        by_who[row[1]] = by_who.get(row[1], 0) + 1
    print(f"Đã ghi {OUT.relative_to(ROOT)}: R1 có {last - 1} dòng", by_who)
    print(f"  chính {len(core)} · nền tảng {min(len(found), T_CAP)}/{len(found)} · paper nền thêm tay {len(SEEDS)}")


if __name__ == "__main__":
    main()
