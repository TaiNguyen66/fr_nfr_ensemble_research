"""Tạo workbook vòng cuối (amendment A-004): 30 bài có PDF + bằng chứng tự động + bảng gap ứng viên.

Chạy (từ gốc repo, sau select_final_papers.py và extract_pdf_evidence.py):
    python scripts/build_final_workbook.py

Ghi literature/review/NCKH_FRNFR_VongCuoi.xlsx gồm:
    0_HuongDan       cách xác nhận và chấm điểm
    VongCuoi         30 bài: thông tin + bằng chứng tự động + cột XÁC NHẬN của người đọc + QA
    GD3_Gap          gap ứng viên (từ gap_candidates.csv) + cột xác nhận / quyết định
    LuongSangLoc     số bài qua từng bước lọc (cho sơ đồ trong bài báo)
    BangChung        câu trích nguyên văn từ PDF cho từng trường
"""

import csv
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
REVIEW = ROOT / "literature" / "review"
SEARCH = ROOT / "literature" / "search"
OUT = REVIEW / "NCKH_FRNFR_VongCuoi.xlsx"
A, B = "Hoang Anh", "Phat Tai"
AUTO = ["dataset", "task", "technique", "split", "leakage_control", "metric", "stat_test", "reproducibility", "imbalance"]

GUIDE = [
    ("Mục đích", "Xác nhận 30 bài vòng cuối và tìm gap. Thông tin tự động (cột xám) do dò từ khóa trong PDF: CÓ THỂ SAI, phải kiểm tra."),
    ("Cách làm", "Mở PDF trong Drive 01_Tai_lieu_tham_khao/Vong_cuoi_30_bai (tên file bắt đầu bằng ID). Đọc Abstract, Method/Evaluation, bảng kết quả (~20–30 phút/bài). "
     "Điền các cột 'XÁC NHẬN' (xanh). Câu trích ở tab BangChung giúp tìm nhanh đoạn cần đọc."),
    ("Bản đọc trước", "Các cột XÁC NHẬN và Q1–Q5 đã được điền sẵn sau khi đọc các đoạn liên quan trong PDF (cột Ghi chú ghi độ chắc chắn). "
     "Người đọc chỉ cần mở PDF, kiểm tra nhanh các ô quan trọng (cách chia, câu trùng, oversampling, kiểm định) — ưu tiên dòng có độ chắc chắn 'thấp' / 'trung bình' — "
     "rồi ghi ĐỒNG Ý hoặc SỬA ở cột cuối. Số liệu gap chỉ tính các dòng đã có người ĐỒNG Ý / SỬA."),
    ("Câu hỏi quan trọng nhất", "Mỗi bài: (1) chia train/test thế nào; (2) câu trùng có thể nằm ở cả train và test không; "
     "(3) oversampling / augmentation làm TRƯỚC hay SAU khi chia; (4) có kiểm định thống kê khi so model không; (5) báo loại F1 nào."),
    ("Bài lạc đề", "Nếu đọc thấy bài không phải phân loại yêu cầu FR/NFR: ghi 'Giữ?' = N và lý do ở Ghi chú. Không cần điền tiếp."),
    ("QA Q1–Q5", "0/1/2 mỗi tiêu chí: Q1 mục tiêu rõ · Q2 dataset mô tả và truy cập được · Q3 đánh giá chặt chẽ (cách chia rõ, không rò rỉ hiển nhiên) · "
     "Q4 kết quả đủ (metric phù hợp, có độ biến thiên) · Q5 nêu hạn chế + tái lập được. Tổng ≥ 6 = tổng hợp chính."),
    ("Gap", "Tab GD3_Gap đã có 7 gap ứng viên với số liệu SƠ BỘ. Sau khi đọc xong, điền cột 'Xác nhận sau khi đọc' bằng số thật (k/N) rồi họp chọn."),
]


def style(ws, widths, fill="DDEBF7"):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "D2"


def short(title, n=50):
    return re.sub(r"[^A-Za-z0-9]+", "_", title)[:n].strip("_")


def main():
    sys.stdout.reconfigure(encoding="utf-8")  # console Windows mặc định cp1252
    final = list(csv.DictReader((REVIEW / "final_set.csv").open(encoding="utf-8-sig")))
    cand = list(csv.DictReader((REVIEW / "final_candidates.csv").open(encoding="utf-8-sig")))
    ev = {r["id"]: r for r in csv.DictReader((REVIEW / "pdf_evidence.csv").open(encoding="utf-8-sig"))}
    gaps = list(csv.reader((REVIEW / "gap_candidates.csv").open(encoding="utf-8-sig")))

    wb = Workbook()
    g = wb.active
    g.title = "0_HuongDan"
    g.append(["Mục", "Nội dung"])
    for row in GUIDE:
        g.append(list(row))
    style(g, [24, 140])
    for r in g.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("VongCuoi")
    meta = ["ID", "Nhóm", "Tiêu đề", "Năm", "Venue", "DOI", "File PDF (Drive)", "Người đọc"]
    auto = [f"[tự động] {k}" for k in AUTO]
    verify = ["XÁC NHẬN: Giữ? (Y/N)", "XÁC NHẬN: Dataset (+version)", "XÁC NHẬN: Nhiệm vụ (binary / multi-class / multi-label)",
              "XÁC NHẬN: Model chính", "XÁC NHẬN: Cách chia dữ liệu", "XÁC NHẬN: Câu trùng có thể ở cả train và test? (Y/N/không rõ)",
              "XÁC NHẬN: Oversampling / augmentation trước khi chia? (Y/N/không dùng)", "XÁC NHẬN: Loại F1 (macro / weighted / micro / theo lớp / chỉ accuracy)",
              "XÁC NHẬN: Kiểm định thống kê (Y/N)", "XÁC NHẬN: Code/data công khai (Y/N)", "XÁC NHẬN: Kết quả tốt nhất (kèm điều kiện)",
              "Q1", "Q2", "Q3", "Q4", "Q5", "QA tổng", "Hạn chế / ý cho đề tài", "Ghi chú",
              "Người kiểm tra: ĐỒNG Ý / SỬA (ghi ô đã sửa)"]
    ws.append(meta + auto + verify)
    # Bản đọc trước (pre_read_verification.csv): điền sẵn vào cột XÁC NHẬN, người đọc chỉ cần kiểm tra
    vpath = REVIEW / "pre_read_verification.csv"
    ver = {r["id"]: r for r in csv.DictReader(vpath.open(encoding="utf-8-sig"))} if vpath.exists() else {}
    vkeys = ["giu", "dataset", "nhiem_vu", "model_chinh", "cach_chia", "cau_trung_ca_train_test", "oversampling_truoc_khi_chia",
             "loai_f1", "kiem_dinh_thong_ke", "code_data", "ket_qua_tot_nhat"]
    pattern = [B, A, B, A, B]  # Phat Tai ~60%, Hoang Anh ~40%
    for i, p in enumerate(sorted(final, key=lambda r: (r["group"], r["id"]))):
        e = ev.get(p["id"], {})
        row = [p["id"], p["group"], p["title"], p["year"], p["venue"][:50], f"https://doi.org/{p['doi']}" if p["doi"] else "",
               f"{p['id']}_{short(p['title'])}.pdf", pattern[i % len(pattern)]]
        v = ver.get(p["id"])
        if v:
            qs = [int(v[q]) if v[q] != "" else None for q in ("q1", "q2", "q3", "q4", "q5")]
            note = " · ".join(x for x in (v["ly_do_loai"], v["ghi_chu"], f"Đọc trước (độ chắc chắn: {v['do_chac_chan']})") if x)
            vals = [v[k] for k in vkeys] + qs + ["", "", note, ""]
        else:
            vals = [""] * len(verify)
        row += [e.get(k, "") for k in AUTO] + vals
        ws.append(row)
        r = ws.max_row
        q = [ws.cell(r, len(meta) + len(auto) + 12 + k).coordinate for k in range(5)]
        ws.cell(r, len(meta) + len(auto) + 17, f"=IF(COUNT({q[0]}:{q[4]})=5,SUM({q[0]}:{q[4]}),\"\")")
    style(ws, [7, 13, 50, 6, 22, 24, 30, 10] + [18] * len(auto) + [16] * len(verify))
    grey, blue = PatternFill("solid", fgColor="EDEDED"), PatternFill("solid", fgColor="E2EFDA")
    for col in range(len(meta) + 1, len(meta) + len(auto) + 1):
        ws.cell(1, col).fill = grey
    for col in range(len(meta) + len(auto) + 1, ws.max_column + 1):
        ws.cell(1, col).fill = blue
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(dv)
    first_q = ws.cell(2, len(meta) + len(auto) + 12).coordinate
    last_q = ws.cell(ws.max_row, len(meta) + len(auto) + 16).coordinate
    dv.add(f"{first_q}:{last_q}")

    gp = wb.create_sheet("GD3_Gap")
    for row in gaps:
        gp.append(row)
    style(gp, [7, 22, 50, 40, 18, 36, 26, 44, 14, 18, 14, 20])
    for r in gp.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    fl = wb.create_sheet("LuongSangLoc")
    status = Counter(r["status"].split(" (")[0] if r["status"].startswith(("NOT", "FINAL")) else r["status"] for r in cand)
    n01 = sum(1 for _ in csv.DictReader((SEARCH / "01_all_records.csv").open(encoding="utf-8-sig")))
    n02 = list(csv.DictReader((SEARCH / "02_after_screening_v1.csv").open(encoding="utf-8-sig")))
    fl.append(["Bước", "Số bài", "Ghi chú"])
    fl.append(["Bản ghi thu về (OpenAlex, Semantic Scholar, arXiv; 8 chuỗi)", n01, "kể cả bản trùng giữa các nguồn"])
    fl.append(["Bản ghi duy nhất", len(n02), "sau khử trùng DOI / tiêu đề"])
    fl.append(["Qua lọc nội dung L1–L4", len(cand), "INCLUDE, có abstract, bình duyệt, đúng phân loại FR/NFR (+ bài nền tảng)"])
    for k, v in sorted(status.items()):
        fl.append([k, v, ""])
    fl.append(["Vòng cuối (có PDF)", len(final), ", ".join(f"{k} {v}" for k, v in Counter(r["group"] for r in final).items())])
    style(fl, [60, 10, 80])

    bc = wb.create_sheet("BangChung")
    keys = ["id", "title"] + [x for k in AUTO for x in (k, f"{k}_evidence")]
    bc.append(keys)
    for p in sorted(final, key=lambda r: r["id"]):
        e = ev.get(p["id"], {})
        bc.append([e.get(k, "") for k in keys])
    style(bc, [7, 40] + [16, 50] * len(AUTO))

    wb.save(OUT)
    who = Counter(ws.cell(r, 8).value for r in range(2, ws.max_row + 1))
    print(f"Đã ghi {OUT.relative_to(ROOT)}: {len(final)} bài, {len(gaps) - 1} gap ứng viên, phân công {dict(who)}")


if __name__ == "__main__":
    main()
